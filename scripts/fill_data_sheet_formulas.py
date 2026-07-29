"""
데이터 시트 T/W/X/Y/Z열 수식 채우기

업무정의:
  T: 사고전환판매플랜코드 = (최초판매플랜코드 K & 무사고년수 S)를 키값으로
     매핑정보 시트 E열에서 동일 키값을 찾아 F열 값을 반환
  W: 현재판매플랜(E) = 경과전환판매플랜코드(M) 여부
  X: 현재판매플랜(E) = 사고전환판매플랜코드(T) 여부
  Y: 경과년수(L) = 무사고년수(S) 여부
  Z: V가 '01'/'02'가 아니고 Y=FALSE이면서 (X=TRUE 이거나 W=TRUE) 이면 '기전환', 아니면 공란
     (원 명세의 두 갈래 조건은 OR(X,W)와 동치)

사용법: python fill_data_sheet_formulas.py <입력.xlsx> <출력.xlsx>
"""
import sys
import openpyxl


def fill(input_path, output_path):
    wb = openpyxl.load_workbook(input_path, data_only=False)
    ws = wb['데이터']
    mp_last = wb['매핑정보'].max_row

    first_row, last_row = 2, ws.max_row

    for r in range(first_row, last_row + 1):
        ws.cell(row=r, column=20).value = (  # T
            f'=IFERROR(INDEX(매핑정보!$F$2:$F${mp_last},'
            f'MATCH(K{r}&S{r},매핑정보!$E$2:$E${mp_last},0)),"")'
        )
        ws.cell(row=r, column=23).value = f'=E{r}=M{r}'  # W
        ws.cell(row=r, column=24).value = f'=E{r}=T{r}'  # X
        ws.cell(row=r, column=25).value = f'=L{r}=S{r}'  # Y
        ws.cell(row=r, column=26).value = (  # Z
            f'=IF(AND(V{r}<>"01",V{r}<>"02",NOT(Y{r}),OR(X{r},W{r})),"기전환","")'
        )

    wb.save(output_path)
    print(f"데이터 시트 T,W,X,Y,Z 완료 (rows {first_row}~{last_row}) -> {output_path}")


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("Usage: python fill_data_sheet_formulas.py <입력.xlsx> <출력.xlsx>")
        sys.exit(1)
    fill(sys.argv[1], sys.argv[2])
