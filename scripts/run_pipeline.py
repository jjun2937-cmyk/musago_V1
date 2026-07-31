"""
무사고전환 데이터/매핑정보 입력 파일 하나로, 아래 3단계를 한 번에 실행하는 통합 실행기.

  1) build_report_v5  : 데이터 시트에 사고전환판매플랜코드 채우고, 매핑정보에 키값
                         채우고, 레이아웃 4종(당월/전체 x 부문별/상품별)을 원본
                         템플릿 서식으로 생성
  2) recalc_util       : LibreOffice로 전체 수식 재계산 (openpyxl은 계산된 값을
                         남기지 않으므로 필수 단계)
  3) split_report_v5   : 데이터 파일(수식 유지)과 결과 파일(레이아웃 값만)로 분리
                         (분리 과정에서 데이터 파일 쪽 수식 캐시값이 openpyxl에
                         의해 지워지므로, 분리 후 데이터 파일을 한 번 더 재계산)

요구사항: LibreOffice(soffice)가 설치되어 있어야 함
  - Ubuntu/Debian: sudo apt-get install libreoffice-calc
  - macOS: brew install --cask libreoffice

사용법:
  python scripts/run_pipeline.py <입력.xlsx> <데이터출력.xlsx> <결과출력.xlsx> [--timeout 초] [--keep-combined]

입력 파일은 시트 2개(데이터 시트: '수금부문명' 헤더 보유 / 매핑정보 시트: 나머지)를
담고 있어야 하며, 시트명은 무엇이든 상관없다.
"""
import argparse
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl

import build_report_v5 as builder
import recalc_util
from split_report_v5 import LAYOUT_SHEET_NAMES, copy_values_only


def run(input_path, data_out_path, result_out_path, timeout=600, keep_combined=False):
    combined_fd, combined_path = tempfile.mkstemp(suffix='.xlsx', prefix='musago_combined_')
    os.close(combined_fd)

    try:
        print(f"[1/3] 입력 파일 분석 및 수식 생성 중... ({input_path})")
        wb = openpyxl.load_workbook(input_path, data_only=False)
        data_ws, mapping_ws = builder.identify_sheets(wb)
        print(f"  - 데이터 시트: '{data_ws.title}' ({data_ws.max_row}행 x {data_ws.max_column}열)")
        print(f"  - 매핑정보 시트: '{mapping_ws.title}' ({mapping_ws.max_row}행 x {mapping_ws.max_column}열)")

        map_info = builder.process_mapping_sheet(mapping_ws)
        cols = builder.process_data_sheet(data_ws)
        builder.fill_data_formulas(
            data_ws, cols, mapping_ws.title,
            map_info['key'], map_info['target'], map_info['last_row'],
        )

        months = builder.get_distinct_values(data_ws, cols['month'])
        products = builder.get_distinct_values(data_ws, cols['product'])
        print(f"  - 계약체결월: {months}")
        print(f"  - 상품명 ({len(products)}종): {products}")

        builder.build_one_layout(wb, data_ws, cols, months, 'dept', builder.DEPT_GROUPS,
                                  'contains', False, '레이아웃1_당월_부문별')
        builder.build_one_layout(wb, data_ws, cols, months, 'dept', builder.DEPT_GROUPS,
                                  'contains', True, '레이아웃2_전체_부문별')
        builder.build_one_layout(wb, data_ws, cols, months, 'product', products,
                                  'exact', False, '레이아웃3_당월_상품별')
        builder.build_one_layout(wb, data_ws, cols, months, 'product', products,
                                  'exact', True, '레이아웃4_전체_상품별')

        wb.save(combined_path)
        wb.close()

        print(f"[2/3] LibreOffice로 전체 수식 재계산 중 (최대 {timeout}초)...")
        result = recalc_util.recalc(combined_path, timeout=timeout)
        if 'error' in result:
            print(f"  ! 재계산 실패: {result['error']}", file=sys.stderr)
            sys.exit(1)
        print(f"  - 상태: {result['status']}, 수식 오류: {result['total_errors']}건")
        if result['total_errors']:
            print("  ! 오류 셀 목록(최대 50개):")
            for c in result['error_cells']:
                print(f"      {c}")

        print(f"[3/3] 데이터 파일 / 결과 파일 분리 중...")
        wb_values = openpyxl.load_workbook(combined_path, data_only=True)
        result_wb = openpyxl.Workbook()
        result_wb.remove(result_wb.active)
        for name in LAYOUT_SHEET_NAMES:
            copy_values_only(wb_values[name], result_wb, name)
        result_wb.save(result_out_path)
        wb_values.close()

        wb_formulas = openpyxl.load_workbook(combined_path, data_only=False)
        for name in LAYOUT_SHEET_NAMES:
            if name in wb_formulas.sheetnames:
                del wb_formulas[name]
        wb_formulas.save(data_out_path)
        wb_formulas.close()

        # *** 중요 *** 위 저장 과정에서 데이터 파일 쪽 수식(사고전환판매플랜코드,
        # 매핑정보 키값)의 캐시된 계산값이 openpyxl에 의해 지워지므로, 다시
        # 한 번 재계산해서 복원한다 (안 하면 그 두 열이 전부 공란으로 보임).
        print("  데이터 파일 재계산 중 (분리 과정에서 지워진 수식 캐시값 복원)...")
        data_recalc = recalc_util.recalc(data_out_path, timeout=timeout)
        if 'error' in data_recalc:
            print(f"  ! 데이터 파일 재계산 실패: {data_recalc['error']}", file=sys.stderr)
            sys.exit(1)
        print(f"  - 데이터 파일 재계산 상태: {data_recalc['status']}, 수식 오류: {data_recalc['total_errors']}건")

        print(f"  - 데이터 파일 -> {data_out_path}")
        print(f"  - 결과 파일   -> {result_out_path}")
        print("완료.")
        return result['total_errors'] == 0 and data_recalc['total_errors'] == 0
    finally:
        if keep_combined:
            print(f"(--keep-combined) 결합 중간 파일 보존됨: {combined_path}")
        else:
            try:
                os.remove(combined_path)
            except OSError:
                pass


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('input', help='입력 파일 (데이터 시트 + 매핑정보 시트)')
    p.add_argument('data_out', help='데이터 파일 출력 경로')
    p.add_argument('result_out', help='결과 파일 출력 경로')
    p.add_argument('--timeout', type=int, default=600, help='LibreOffice 재계산 제한시간(초), 기본 600')
    p.add_argument('--keep-combined', action='store_true', help='중간 결합 파일을 지우지 않고 보존')
    args = p.parse_args()

    ok = run(args.input, args.data_out, args.result_out,
             timeout=args.timeout, keep_combined=args.keep_combined)
    sys.exit(0 if ok else 2)


if __name__ == '__main__':
    main()
