"""
무사고전환 데이터/매핑정보 -> 데이터(T~Z 추가)/매핑정보(키값 추가)/결과레이아웃/추천요약
한 파일로 생성하는 통합 스크립트.

입력: 시트명과 무관하게, 헤더 내용으로 "데이터 시트"(수금부문명 헤더 보유)와
      "매핑정보 시트"(그 외)를 자동 식별한다.

사용법: python build_report.py <입력.xlsx> <출력.xlsx>
"""
import sys
import re
from collections import Counter
import openpyxl
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# 헤더 탐색 유틸
# ---------------------------------------------------------------------------

def header_map(ws):
    """1행 헤더 텍스트 -> 1-indexed 컬럼 번호. 중복 헤더는 첫 번째 위치만 남긴다."""
    m = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v is not None and v not in m:
            m[v] = c
    return m


def find_column(ws, hmap, name, fallback_anchor=None, fallback_expected=None,
                 rename_to=None):
    """
    이름(name)으로 헤더를 찾는다. 없으면 fallback_anchor 헤더 바로 오른쪽 열이
    fallback_expected 헤더인지 확인하고 그 열을 사용한다.
    rename_to가 주어지면 실제 사용하는 열의 1행 헤더 값을 그 이름으로 바꿔쓴다.
    반환값: 1-indexed 컬럼 번호
    """
    if name in hmap:
        col = hmap[name]
    else:
        if fallback_anchor not in hmap:
            raise ValueError(
                f"'{name}' 헤더가 없고, 대체 기준 헤더 '{fallback_anchor}'도 찾을 수 없습니다."
            )
        anchor_col = hmap[fallback_anchor]
        col = anchor_col + 1
        actual = ws.cell(row=1, column=col).value
        if actual != fallback_expected:
            raise ValueError(
                f"'{fallback_anchor}' 바로 오른쪽 열의 헤더가 '{fallback_expected}'일 것으로 "
                f"예상했으나 실제로는 '{actual}' 입니다. (열 {col})"
            )
    if rename_to and ws.cell(row=1, column=col).value != rename_to:
        ws.cell(row=1, column=col).value = rename_to
    return col


def identify_sheets(wb):
    data_ws = mapping_ws = None
    for ws in wb.worksheets:
        hmap = header_map(ws)
        if '수금부문명' in hmap:
            data_ws = ws
        else:
            mapping_ws = ws
    if data_ws is None or mapping_ws is None:
        raise ValueError("데이터 시트('수금부문명' 헤더 보유) 또는 매핑정보 시트를 찾지 못했습니다.")
    return data_ws, mapping_ws


def find_or_append_column(ws, hmap, name):
    """헤더가 이미 있으면 그 열 번호를, 없으면 맨 우측에 새로 만들고 그 열 번호를 반환."""
    if name in hmap:
        return hmap[name]
    col = ws.max_column + 1
    ws.cell(row=1, column=col).value = name
    hmap[name] = col
    return col


# ---------------------------------------------------------------------------
# 1) 데이터 시트: 필요 컬럼 확정 + T/W/X/Y/Z 채우기
# ---------------------------------------------------------------------------

def process_data_sheet(data_ws):
    hmap = header_map(data_ws)

    col_current = find_column(
        data_ws, hmap, '현재판매플랜코드',
        fallback_anchor='상품명', fallback_expected='판매플랜코드',
        rename_to='현재판매플랜코드',
    )
    col_initial = find_column(
        data_ws, hmap, '최초판매플랜코드',
        fallback_anchor='적용시작일시', fallback_expected='판매플랜코드',
        rename_to='최초판매플랜코드',
    )
    col_elapsed_plan = find_column(
        data_ws, hmap, '경과전환판매플랜코드',
        fallback_anchor='세부플랜코드', fallback_expected='전환판매플랜코드',
        rename_to='경과전환판매플랜코드',
    )
    # 이름 변경이 있었으니 헤더맵 갱신
    hmap = header_map(data_ws)

    col_elapsed_years = hmap['경과년수']
    col_accident_free_years = hmap['무사고년수']
    col_accident_code = hmap['사고구분코드']
    col_process_code = hmap['전환처리구분코드']
    col_dept = hmap['수금부문명']
    col_month = hmap['계약체결월']

    col_T = find_or_append_column(data_ws, hmap, '사고전환판매플랜코드')
    col_W = find_or_append_column(data_ws, hmap, '현재판매플랜=경과전환판매플랜코드')
    col_X = find_or_append_column(data_ws, hmap, '현재판매플랜=사고전환판매플랜코드')
    col_Y = find_or_append_column(data_ws, hmap, '경과년수=무사고년수')
    col_Z = find_or_append_column(data_ws, hmap, '기전환완료여부')

    return {
        'current': col_current, 'initial': col_initial, 'elapsed_plan': col_elapsed_plan,
        'elapsed_years': col_elapsed_years, 'accident_free_years': col_accident_free_years,
        'accident_code': col_accident_code, 'process_code': col_process_code,
        'dept': col_dept, 'month': col_month,
        'T': col_T, 'W': col_W, 'X': col_X, 'Y': col_Y, 'Z': col_Z,
    }


def fill_data_formulas(data_ws, cols, mapping_sheet_name, map_key_col, map_target_col, map_last_row):
    L = get_column_letter
    E = L(cols['current'])
    K = L(cols['initial'])
    Lc = L(cols['elapsed_years'])
    M = L(cols['elapsed_plan'])
    S = L(cols['accident_free_years'])
    V = L(cols['accident_code'])
    T = L(cols['T'])
    W = L(cols['W'])
    X = L(cols['X'])
    Y = L(cols['Y'])
    Z = L(cols['Z'])

    map_key = L(map_key_col)
    map_target = L(map_target_col)

    first_row, last_row = 2, data_ws.max_row
    for r in range(first_row, last_row + 1):
        data_ws.cell(row=r, column=cols['T']).value = (
            f'=IFERROR(INDEX({mapping_sheet_name}!${map_target}$2:${map_target}${map_last_row},'
            f'MATCH({K}{r}&{S}{r},{mapping_sheet_name}!${map_key}$2:${map_key}${map_last_row},0)),"")'
        )
        data_ws.cell(row=r, column=cols['W']).value = f'={E}{r}={M}{r}'
        data_ws.cell(row=r, column=cols['X']).value = f'={E}{r}={T}{r}'
        data_ws.cell(row=r, column=cols['Y']).value = f'={Lc}{r}={S}{r}'
        data_ws.cell(row=r, column=cols['Z']).value = (
            f'=IF(AND({V}{r}<>"01",{V}{r}<>"02",NOT({Y}{r}),OR({X}{r},{W}{r})),"기전환","")'
        )


# ---------------------------------------------------------------------------
# 2) 매핑정보 시트: 키값 컬럼 생성
# ---------------------------------------------------------------------------

def process_mapping_sheet(mapping_ws):
    hmap = header_map(mapping_ws)
    col_plan = hmap['판매플랜코드']
    col_years = hmap['경과년수']
    col_target = hmap['전환판매플랜코드']
    col_key = find_or_append_column(mapping_ws, hmap, '키값')

    L = get_column_letter
    plan_l, years_l, key_l = L(col_plan), L(col_years), L(col_key)
    last_row = mapping_ws.max_row
    for r in range(2, last_row + 1):
        mapping_ws.cell(row=r, column=col_key).value = f'={plan_l}{r}&{years_l}{r}'

    return {'key': col_key, 'target': col_target, 'last_row': last_row}



# ---------------------------------------------------------------------------
# 3) 결과레이아웃 시트: 월별 동적 블록 (개인/전략/신사업/법인 + 합계)
# ---------------------------------------------------------------------------

GROUPS = [('개인', '개인'), ('전략', '전략'), ('신사업', '신사업'), ('법인', '법인')]
CODES = [('O', '001'), ('P', '002'), ('Q', '003'), ('R', '004'),
         ('S', '005'), ('T', '006'), ('U', '007')]

LAYOUT_HEADERS = [
    '구분', '전체계약건수', '사고有', '전환대상(A-B)', '기존전환완료', '전환미완료(C-D)',
    '기존미전환', '당월전환대상(E-F)', '전환완료', '완료율(%)', '전환미완료(G-H)', '미완료율(%)',
    '현장활동확인 합계', '전체比(%)', '001 연락두절 등', '002 사고고지', '003 전환의사無',
    '004 압류계약', '005 전환예정', '006 기타', '007 ARS거부', '현장활동미확인', '전체比(%)',
]


def get_distinct_months(data_ws, month_col):
    months = set()
    for r in range(2, data_ws.max_row + 1):
        v = data_ws.cell(row=r, column=month_col).value
        if v not in (None, ''):
            months.add(str(v))
    return sorted(months)


def build_layout_sheet(wb, data_ws, cols, months, sheet_name='결과레이아웃'):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    L = get_column_letter
    data_name = f"'{data_ws.title}'" if ' ' in data_ws.title else data_ws.title
    last_row = data_ws.max_row
    dept_l = L(cols['dept'])
    month_l = L(cols['month'])
    v_l = L(cols['accident_code'])
    z_l = L(cols['Z'])
    y_l = L(cols['Y'])
    x_l = L(cols['X'])
    u_l = L(cols['process_code'])

    for i, h in enumerate(LAYOUT_HEADERS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.value = h
        cell.font = openpyxl.styles.Font(bold=True)

    blocks = []  # (month, total_row, [group_rows...])
    row = 2
    for month in months:
        total_row = row
        group_first = row + 1
        group_last = row + 4
        blocks.append((month, total_row, list(range(group_first, group_last + 1))))

        ws.cell(row=total_row, column=1).value = f"{month}월"
        ws.cell(row=total_row, column=1).font = openpyxl.styles.Font(bold=True)
        for col_letter in ['B', 'C', 'E', 'G', 'I']:
            ws[f'{col_letter}{total_row}'] = f'=SUM({col_letter}{group_first}:{col_letter}{group_last})'
        ws[f'D{total_row}'] = f'=B{total_row}-C{total_row}'
        ws[f'F{total_row}'] = f'=D{total_row}-E{total_row}'
        ws[f'H{total_row}'] = f'=F{total_row}-G{total_row}'
        ws[f'J{total_row}'] = f'=IFERROR(I{total_row}/H{total_row}*100,"")'
        ws[f'K{total_row}'] = f'=H{total_row}-I{total_row}'
        ws[f'L{total_row}'] = f'=IFERROR(K{total_row}/H{total_row}*100,"")'
        ws[f'M{total_row}'] = f'=SUM(O{total_row}:U{total_row})'
        ws[f'N{total_row}'] = f'=IFERROR(M{total_row}/H{total_row}*100,"")'
        for col_letter, _ in CODES:
            ws[f'{col_letter}{total_row}'] = f'=SUM({col_letter}{group_first}:{col_letter}{group_last})'
        ws[f'V{total_row}'] = f'=K{total_row}-M{total_row}'
        ws[f'W{total_row}'] = f'=IFERROR(V{total_row}/H{total_row}*100,"")'

        dept_rng = f"{data_name}!${dept_l}$2:${dept_l}${last_row}"
        month_rng = f"{data_name}!${month_l}$2:${month_l}${last_row}"
        v_rng = f"{data_name}!${v_l}$2:${v_l}${last_row}"
        z_rng = f"{data_name}!${z_l}$2:${z_l}${last_row}"
        y_rng = f"{data_name}!${y_l}$2:${y_l}${last_row}"
        x_rng = f"{data_name}!${x_l}$2:${x_l}${last_row}"
        u_rng = f"{data_name}!${u_l}$2:${u_l}${last_row}"

        for gi, (label, keyword) in enumerate(GROUPS):
            r = group_first + gi
            ws.cell(row=r, column=1).value = label

            base = f'({month_rng}="{month}")*ISNUMBER(SEARCH("{keyword}",{dept_rng}))'

            ws[f'B{r}'] = f'=SUMPRODUCT({base})'
            ws[f'C{r}'] = f'=SUMPRODUCT({base}*(({v_rng}="01")+({v_rng}="02")))'
            ws[f'D{r}'] = f'=B{r}-C{r}'
            ws[f'E{r}'] = f'=SUMPRODUCT({base}*({z_rng}="기전환"))'
            ws[f'F{r}'] = f'=D{r}-E{r}'
            ws[f'G{r}'] = (
                f'=SUMPRODUCT({base}*({v_rng}<>"01")*({v_rng}<>"02")*'
                f'({z_rng}<>"기전환")*({y_rng}=FALSE))'
            )
            ws[f'H{r}'] = f'=F{r}-G{r}'
            ws[f'I{r}'] = (
                f'=SUMPRODUCT({base}*({v_rng}<>"01")*({v_rng}<>"02")*'
                f'({z_rng}<>"기전환")*({y_rng}=TRUE)*({x_rng}=TRUE))'
            )
            ws[f'J{r}'] = f'=IFERROR(I{r}/H{r}*100,"")'
            ws[f'K{r}'] = f'=H{r}-I{r}'
            ws[f'L{r}'] = f'=IFERROR(K{r}/H{r}*100,"")'
            for col_letter, code in CODES:
                ws[f'{col_letter}{r}'] = (
                    f'=SUMPRODUCT({base}*({v_rng}<>"01")*({v_rng}<>"02")*({z_rng}<>"기전환")*'
                    f'({y_rng}=TRUE)*({x_rng}=FALSE)*({u_rng}="{code}"))'
                )
            ws[f'M{r}'] = f'=SUM(O{r}:U{r})'
            ws[f'N{r}'] = f'=IFERROR(M{r}/H{r}*100,"")'
            ws[f'V{r}'] = f'=K{r}-M{r}'
            ws[f'W{r}'] = f'=IFERROR(V{r}/H{r}*100,"")'

        row = group_last + 2  # 그룹행 4개 다음 빈 줄 1개

    return ws, blocks


# ---------------------------------------------------------------------------
# 4) 추천 보고양식 시트: 결과레이아웃 값을 참조해 월x부문 한 화면 요약
# ---------------------------------------------------------------------------

def build_summary_sheet(wb, layout_ws, blocks, sheet_name='추천요약'):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    layout_name = f"'{layout_ws.title}'" if ' ' in layout_ws.title else layout_ws.title

    headers = ['월', '구분', '전체계약건수', '전환대상', '기전환완료', '당월전환대상',
               '전환완료', '완료율(%)', '전환미완료', '현장활동확인', '현장활동미확인']
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.value = h
        c.font = openpyxl.styles.Font(bold=True)
        c.fill = openpyxl.styles.PatternFill('solid', fgColor='DDEBF7')

    r = 2
    for month, total_row, group_rows in blocks:
        rows_to_show = [(f'{month}월 합계', total_row)] + [
            (layout_ws.cell(row=gr, column=1).value, gr) for gr in group_rows
        ]
        for label, src_row in rows_to_show:
            ws.cell(row=r, column=1).value = f"{month}월" if src_row == total_row else ''
            ws.cell(row=r, column=2).value = label
            ws.cell(row=r, column=3).value = f'={layout_name}!B{src_row}'
            ws.cell(row=r, column=4).value = f'={layout_name}!D{src_row}'
            ws.cell(row=r, column=5).value = f'={layout_name}!E{src_row}'
            ws.cell(row=r, column=6).value = f'={layout_name}!H{src_row}'
            ws.cell(row=r, column=7).value = f'={layout_name}!I{src_row}'
            ws.cell(row=r, column=8).value = f'={layout_name}!J{src_row}'
            ws.cell(row=r, column=9).value = f'={layout_name}!K{src_row}'
            ws.cell(row=r, column=10).value = f'={layout_name}!M{src_row}'
            ws.cell(row=r, column=11).value = f'={layout_name}!V{src_row}'
            if src_row == total_row:
                for col in range(1, 12):
                    ws.cell(row=r, column=col).font = openpyxl.styles.Font(bold=True)
            r += 1

    # 완료율(%) 열에 3색 스케일 조건부서식 (DataBarRule은 x14 확장을 써서 LibreOffice
    # 재계산 시 제거되므로, 호환성 좋은 ColorScaleRule 사용)
    from openpyxl.formatting.rule import ColorScaleRule
    rule = ColorScaleRule(
        start_type='num', start_value=0, start_color='F8696B',
        mid_type='num', mid_value=50, mid_color='FFEB84',
        end_type='num', end_value=100, end_color='63BE7B',
    )
    ws.conditional_formatting.add(f'H2:H{r-1}', rule)

    widths = [8, 12, 14, 10, 12, 12, 10, 10, 10, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    return ws


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("Usage: python build_report.py <입력.xlsx> <출력.xlsx>")
        sys.exit(1)

    in_path, out_path = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(in_path, data_only=False)
    data_ws, mapping_ws = identify_sheets(wb)
    print(f"데이터 시트 식별: '{data_ws.title}' ({data_ws.max_row}행 x {data_ws.max_column}열)")
    print(f"매핑정보 시트 식별: '{mapping_ws.title}' ({mapping_ws.max_row}행 x {mapping_ws.max_column}열)")

    map_info = process_mapping_sheet(mapping_ws)
    cols = process_data_sheet(data_ws)
    fill_data_formulas(
        data_ws, cols, mapping_ws.title,
        map_info['key'], map_info['target'], map_info['last_row'],
    )

    months = get_distinct_months(data_ws, cols['month'])
    print("발견된 계약체결월:", months)
    layout_ws, blocks = build_layout_sheet(wb, data_ws, cols, months)
    build_summary_sheet(wb, layout_ws, blocks)

    wb.save(out_path)
    print(f"전체 완료 -> {out_path}")
