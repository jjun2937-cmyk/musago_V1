"""
무사고전환 데이터/매핑정보 -> 데이터(사고전환판매플랜코드 추가)/매핑정보(키값 추가)/
결과레이아웃(v3 신규 지표)/추천요약 한 파일로 생성하는 통합 스크립트 (v3).

v2(build_report.py)와의 차이 - 결과레이아웃 지표 정의를 통째로 교체:
  - "당월" 그룹과 "기존" 그룹을 사고구분코드('01'=당사사고 / '02'=타사사고)로 가른다.
    - 사고구분코드 <> '01' 인 계약 -> "당월" 버킷 (당월전환대상/완료/미완료)
    - 사고구분코드 = '01' 인 계약 -> "기존" 버킷 (기존전환완료/미완료)
  - 당월전환완료 = (사고구분코드<>'01') 중 현재판매플랜코드 = 경과전환판매플랜코드
    (경과년수=무사고년수 여부는 더 이상 따지지 않는다 - 다르더라도 플랜이 이미
    경과전환판매플랜코드와 같으면 완료로 인정)
  - 기존전환완료 = (사고구분코드='01') 중 현재판매플랜코드 = 사고전환판매플랜코드
    (사고전환판매플랜코드는 최초판매플랜코드+무사고년수로 매핑정보를 조회한 값)

입력: 시트명과 무관하게, 헤더 내용으로 "데이터 시트"(수금부문명 헤더 보유)와
      "매핑정보 시트"(그 외)를 자동 식별한다.

사용법: python build_report_v3.py <입력.xlsx> <출력.xlsx>
"""
import sys
import openpyxl
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# 헤더 탐색 유틸 (v2와 동일)
# ---------------------------------------------------------------------------

def header_map(ws):
    m = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v is not None and v not in m:
            m[v] = c
    return m


def find_column(ws, hmap, name, fallback_anchor=None, fallback_expected=None, rename_to=None):
    if name in hmap:
        col = hmap[name]
    else:
        if fallback_anchor not in hmap:
            raise ValueError(f"'{name}' 헤더가 없고, 대체 기준 헤더 '{fallback_anchor}'도 찾을 수 없습니다.")
        anchor_col = hmap[fallback_anchor]
        col = anchor_col + 1
        actual = ws.cell(row=1, column=col).value
        if actual != fallback_expected:
            raise ValueError(
                f"'{fallback_anchor}' 바로 오른쪽 열의 헤더가 '{fallback_expected}'일 것으로 "
                f"예상했으나 실제로는 '{actual}' 입니다. (열 {col})"
            )
    if rename_to and ws.cell(row=1, column=col).value != rename_to:
        ws.cell(row=1, column=col).value = rename_to
    return col


def identify_sheets(wb):
    data_ws = mapping_ws = None
    for ws in wb.worksheets:
        hmap = header_map(ws)
        if '수금부문명' in hmap:
            data_ws = ws
        else:
            mapping_ws = ws
    if data_ws is None or mapping_ws is None:
        raise ValueError("데이터 시트('수금부문명' 헤더 보유) 또는 매핑정보 시트를 찾지 못했습니다.")
    return data_ws, mapping_ws


def find_or_append_column(ws, hmap, name):
    if name in hmap:
        return hmap[name]
    col = ws.max_column + 1
    ws.cell(row=1, column=col).value = name
    hmap[name] = col
    return col


# ---------------------------------------------------------------------------
# 1) 데이터 시트: 필요 컬럼 확정 + 사고전환판매플랜코드(T) 채우기
# ---------------------------------------------------------------------------

def process_data_sheet(data_ws):
    hmap = header_map(data_ws)

    col_current = find_column(
        data_ws, hmap, '현재판매플랜코드',
        fallback_anchor='상품명', fallback_expected='판매플랜코드',
        rename_to='현재판매플랜코드',
    )
    col_initial = find_column(
        data_ws, hmap, '최초판매플랜코드',
        fallback_anchor='적용시작일시', fallback_expected='판매플랜코드',
        rename_to='최초판매플랜코드',
    )
    col_elapsed_plan = find_column(
        data_ws, hmap, '경과전환판매플랜코드',
        fallback_anchor='세부플랜코드', fallback_expected='전환판매플랜코드',
        rename_to='경과전환판매플랜코드',
    )
    hmap = header_map(data_ws)  # 이름 변경 반영

    col_accident_free_years = hmap['무사고년수']
    col_accident_code = hmap['사고구분코드']
    col_process_code = hmap['전환처리구분코드']
    col_dept = hmap['수금부문명']
    col_month = hmap['계약체결월']

    col_T = find_or_append_column(data_ws, hmap, '사고전환판매플랜코드')

    return {
        'current': col_current, 'initial': col_initial, 'elapsed_plan': col_elapsed_plan,
        'accident_free_years': col_accident_free_years, 'accident_code': col_accident_code,
        'process_code': col_process_code, 'dept': col_dept, 'month': col_month, 'T': col_T,
    }


def fill_data_formulas(data_ws, cols, mapping_sheet_name, map_key_col, map_target_col, map_last_row):
    L = get_column_letter
    K = L(cols['initial'])
    S = L(cols['accident_free_years'])
    map_key = L(map_key_col)
    map_target = L(map_target_col)

    first_row, last_row = 2, data_ws.max_row
    for r in range(first_row, last_row + 1):
        data_ws.cell(row=r, column=cols['T']).value = (
            f'=IFERROR(INDEX({mapping_sheet_name}!${map_target}$2:${map_target}${map_last_row},'
            f'MATCH({K}{r}&{S}{r},{mapping_sheet_name}!${map_key}$2:${map_key}${map_last_row},0)),"")'
        )


# ---------------------------------------------------------------------------
# 2) 매핑정보 시트: 키값 컬럼 생성 (v2와 동일)
# ---------------------------------------------------------------------------

def process_mapping_sheet(mapping_ws):
    hmap = header_map(mapping_ws)
    col_plan = hmap['판매플랜코드']
    col_years = hmap['경과년수']
    col_target = hmap['전환판매플랜코드']
    col_key = find_or_append_column(mapping_ws, hmap, '키값')

    L = get_column_letter
    plan_l, years_l = L(col_plan), L(col_years)
    last_row = mapping_ws.max_row
    for r in range(2, last_row + 1):
        mapping_ws.cell(row=r, column=col_key).value = f'={plan_l}{r}&{years_l}{r}'

    return {'key': col_key, 'target': col_target, 'last_row': last_row}


# ---------------------------------------------------------------------------
# 3) 결과레이아웃 시트 (v3): 월별 동적 블록
# ---------------------------------------------------------------------------

GROUPS = [('개인', '개인'), ('전략', '전략'), ('신사업', '신사업'), ('법인', '법인')]
CODES = [('L', '001', '연락두절 등'), ('M', '002', '사고有(타사)'), ('N', '003', '전환의사無'),
         ('O', '004', '압류계약'), ('P', '005', '전환예정'), ('Q', '006', '기타'),
         ('R', '007', 'ARS거부')]

LAYOUT_HEADERS = {
    'B': '구분', 'C': '유지계약(A)', 'D': '사고有(당사)(B)', 'E': '당월전환대상(C=A-B)',
    'F': '당월전환완료', 'G': '완료율(%)', 'H': '당월전환미완료', 'I': '미완료율(%)',
    'J': '현장활동확인', 'K': '전체比(%)',
    'L': '001 연락두절 등', 'M': '002 사고有(타사)', 'N': '003 전환의사無',
    'O': '004 압류계약', 'P': '005 전환예정', 'Q': '006 기타', 'R': '007 ARS거부',
    'S': '현장활동미확인', 'T': '전체比(%)',
    'U': '기존전환완료', 'V': '유지계약比(%)', 'W': '기존전환미완료', 'X': '유지계약比(%)',
}


def get_distinct_months(data_ws, month_col):
    months = set()
    for r in range(2, data_ws.max_row + 1):
        v = data_ws.cell(row=r, column=month_col).value
        if v not in (None, ''):
            months.add(str(v))
    return sorted(months)


def build_layout_sheet(wb, data_ws, cols, months, sheet_name='결과레이아웃'):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    L = get_column_letter
    data_name = f"'{data_ws.title}'" if ' ' in data_ws.title else data_ws.title
    last_row = data_ws.max_row
    dept_l = L(cols['dept'])
    month_l = L(cols['month'])
    v_l = L(cols['accident_code'])
    e_l = L(cols['current'])
    m_l = L(cols['elapsed_plan'])
    t_l = L(cols['T'])
    u_l = L(cols['process_code'])

    for col_letter, label in LAYOUT_HEADERS.items():
        cell = ws[f'{col_letter}1']
        cell.value = label
        cell.font = openpyxl.styles.Font(bold=True)

    dept_rng = f"{data_name}!${dept_l}$2:${dept_l}${last_row}"
    month_rng = f"{data_name}!${month_l}$2:${month_l}${last_row}"
    v_rng = f"{data_name}!${v_l}$2:${v_l}${last_row}"
    e_rng = f"{data_name}!${e_l}$2:${e_l}${last_row}"
    m_rng = f"{data_name}!${m_l}$2:${m_l}${last_row}"
    t_rng = f"{data_name}!${t_l}$2:${t_l}${last_row}"
    u_rng = f"{data_name}!${u_l}$2:${u_l}${last_row}"

    def write_row(r, month, extra_filter=None):
        """extra_filter: 그룹 필터(부문 키워드 조건 문자열) 또는 None(월 전체)."""
        base = f'({month_rng}="{month}")'
        if extra_filter:
            base += f'*{extra_filter}'

        ws[f'C{r}'] = f'=SUMPRODUCT({base})'
        ws[f'D{r}'] = f'=SUMPRODUCT({base}*({v_rng}="01"))'
        ws[f'E{r}'] = f'=C{r}-D{r}'
        ws[f'F{r}'] = f'=SUMPRODUCT({base}*({v_rng}<>"01")*({e_rng}={m_rng}))'
        ws[f'G{r}'] = f'=IFERROR(F{r}/E{r}*100,"")'
        ws[f'H{r}'] = f'=E{r}-F{r}'
        ws[f'I{r}'] = f'=IFERROR(H{r}/E{r}*100,"")'
        for col_letter, code, _label in CODES:
            ws[f'{col_letter}{r}'] = (
                f'=SUMPRODUCT({base}*({v_rng}<>"01")*({e_rng}<>{m_rng})*({u_rng}="{code}"))'
            )
        ws[f'J{r}'] = f'=SUM(L{r}:R{r})'
        ws[f'K{r}'] = f'=IFERROR(J{r}/E{r}*100,"")'
        ws[f'S{r}'] = f'=H{r}-J{r}'
        ws[f'T{r}'] = f'=IFERROR(S{r}/E{r}*100,"")'
        ws[f'U{r}'] = f'=SUMPRODUCT({base}*({v_rng}="01")*({e_rng}={t_rng}))'
        ws[f'V{r}'] = f'=IFERROR(U{r}/C{r}*100,"")'
        ws[f'W{r}'] = f'=D{r}-U{r}'
        ws[f'X{r}'] = f'=IFERROR(W{r}/C{r}*100,"")'

    blocks = []  # (month, total_row, [group_rows...], unclassified_row)
    row = 2
    for month in months:
        total_row = row
        group_first = row + 1
        group_last = row + 4
        unclassified_row = group_last + 1
        blocks.append((month, total_row, list(range(group_first, group_last + 1)), unclassified_row))

        ws.cell(row=total_row, column=2).value = f"{month}월"
        ws.cell(row=total_row, column=2).font = openpyxl.styles.Font(bold=True)
        write_row(total_row, month)  # 그룹 합산이 아니라 월 전체를 독립 재집계

        for gi, (label, keyword) in enumerate(GROUPS):
            r = group_first + gi
            ws.cell(row=r, column=2).value = label
            write_row(r, month, extra_filter=f'ISNUMBER(SEARCH("{keyword}",{dept_rng}))')

        # 미분류 행: 그룹 키워드 어디에도 안 걸리는 수금부문명 건수 (데이터 품질 점검용)
        ws.cell(row=unclassified_row, column=2).value = '미분류(그룹매칭 안됨)'
        not_any_kw = ''.join(
            f'*(ISNUMBER(SEARCH("{kw}",{dept_rng}))=FALSE)' for _, kw in GROUPS
        )
        ws[f'C{unclassified_row}'] = f'=SUMPRODUCT(({month_rng}="{month}"){not_any_kw})'

        row = unclassified_row + 2  # 미분류행 다음 빈 줄 1개

    return ws, blocks


# ---------------------------------------------------------------------------
# 4) 추천 보고양식 시트: 결과레이아웃 값을 참조해 월x부문 한 화면 요약
# ---------------------------------------------------------------------------

def build_summary_sheet(wb, layout_ws, blocks, sheet_name='추천요약'):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    layout_name = f"'{layout_ws.title}'" if ' ' in layout_ws.title else layout_ws.title

    headers = ['월', '구분', '유지계약', '당월전환대상', '당월전환완료', '완료율(%)',
               '당월전환미완료', '현장활동확인', '현장활동미확인', '기존전환완료', '기존전환미완료']
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.value = h
        c.font = openpyxl.styles.Font(bold=True)
        c.fill = openpyxl.styles.PatternFill('solid', fgColor='DDEBF7')

    r = 2
    for month, total_row, group_rows, unclassified_row in blocks:
        rows_to_show = [(f'{month}월 합계', total_row)] + [
            (layout_ws.cell(row=gr, column=2).value, gr) for gr in group_rows
        ]
        for label, src_row in rows_to_show:
            ws.cell(row=r, column=1).value = f"{month}월" if src_row == total_row else ''
            ws.cell(row=r, column=2).value = label
            ws.cell(row=r, column=3).value = f'={layout_name}!C{src_row}'
            ws.cell(row=r, column=4).value = f'={layout_name}!E{src_row}'
            ws.cell(row=r, column=5).value = f'={layout_name}!F{src_row}'
            ws.cell(row=r, column=6).value = f'={layout_name}!G{src_row}'
            ws.cell(row=r, column=7).value = f'={layout_name}!H{src_row}'
            ws.cell(row=r, column=8).value = f'={layout_name}!J{src_row}'
            ws.cell(row=r, column=9).value = f'={layout_name}!S{src_row}'
            ws.cell(row=r, column=10).value = f'={layout_name}!U{src_row}'
            ws.cell(row=r, column=11).value = f'={layout_name}!W{src_row}'
            if src_row == total_row:
                for col in range(1, 12):
                    ws.cell(row=r, column=col).font = openpyxl.styles.Font(bold=True)
            r += 1
        ws.cell(row=r, column=2).value = '미분류(그룹매칭 안됨)'
        ws.cell(row=r, column=3).value = f'={layout_name}!C{unclassified_row}'
        ws.cell(row=r, column=2).font = openpyxl.styles.Font(italic=True, color='C00000')
        ws.cell(row=r, column=3).font = openpyxl.styles.Font(italic=True, color='C00000')
        r += 1

    from openpyxl.formatting.rule import ColorScaleRule
    rule = ColorScaleRule(
        start_type='num', start_value=0, start_color='F8696B',
        mid_type='num', mid_value=50, mid_color='FFEB84',
        end_type='num', end_value=100, end_color='63BE7B',
    )
    ws.conditional_formatting.add(f'F2:F{r-1}', rule)

    widths = [8, 20, 12, 14, 14, 10, 14, 14, 14, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    return ws


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("Usage: python build_report_v3.py <입력.xlsx> <출력.xlsx>")
        sys.exit(1)

    in_path, out_path = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(in_path, data_only=False)
    data_ws, mapping_ws = identify_sheets(wb)
    print(f"데이터 시트 식별: '{data_ws.title}' ({data_ws.max_row}행 x {data_ws.max_column}열)")
    print(f"매핑정보 시트 식별: '{mapping_ws.title}' ({mapping_ws.max_row}행 x {mapping_ws.max_column}열)")

    map_info = process_mapping_sheet(mapping_ws)
    cols = process_data_sheet(data_ws)
    fill_data_formulas(
        data_ws, cols, mapping_ws.title,
        map_info['key'], map_info['target'], map_info['last_row'],
    )

    months = get_distinct_months(data_ws, cols['month'])
    print("발견된 계약체결월:", months)
    layout_ws, blocks = build_layout_sheet(wb, data_ws, cols, months)
    build_summary_sheet(wb, layout_ws, blocks)

    wb.save(out_path)
    print(f"전체 완료 -> {out_path}")
