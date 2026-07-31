"""
build_report_v4.py로 만들고 recalc.py로 재계산까지 끝낸 결합 파일을
"데이터 파일"(데이터+매핑정보, 수식 유지)과 "결과 파일"(레이아웃 4종, 값만)로 분리.

레이아웃 시트의 수식은 데이터 시트를 참조하는데, 결과 파일에는 데이터 시트가 없으므로
그대로 두면 #REF! 가 난다. 그래서 결과 파일 쪽은 재계산된 값만 복사해서 넣는다.

사용법: python split_report_v4.py <재계산된_결합.xlsx> <데이터출력.xlsx> <결과출력.xlsx>
"""
import sys
import openpyxl

LAYOUT_SHEET_NAMES = [
    '레이아웃1_당월_부문별', '레이아웃2_전체_부문별',
    '레이아웃3_당월_상품별', '레이아웃4_전체_상품별',
]


def copy_values_only(src_ws, dst_wb, sheet_name):
    dst_ws = dst_wb.create_sheet(sheet_name)
    for row in src_ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            new_cell = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font = cell.font.copy()
                new_cell.fill = cell.fill.copy()
                new_cell.number_format = cell.number_format
    for col_letter, dim in src_ws.column_dimensions.items():
        if dim.width:
            dst_ws.column_dimensions[col_letter].width = dim.width
    return dst_ws


if __name__ == '__main__':
    if len(sys.argv) != 4:
        print("Usage: python split_report_v4.py <재계산된_결합.xlsx> <데이터출력.xlsx> <결과출력.xlsx>")
        sys.exit(1)

    combined_path, data_out_path, result_out_path = sys.argv[1], sys.argv[2], sys.argv[3]

    # 1) 결과 파일: 재계산된 값(data_only=True)만 읽어서 새 워크북에 값으로 복사
    wb_values = openpyxl.load_workbook(combined_path, data_only=True)
    result_wb = openpyxl.Workbook()
    result_wb.remove(result_wb.active)
    for name in LAYOUT_SHEET_NAMES:
        copy_values_only(wb_values[name], result_wb, name)
    result_wb.save(result_out_path)
    print(f"결과 파일(값만, 레이아웃 {len(LAYOUT_SHEET_NAMES)}종) 저장 완료 -> {result_out_path}")

    # 2) 데이터 파일: 수식 유지된 원본(data_only=False)에서 레이아웃 시트만 제거하고 저장
    wb_formulas = openpyxl.load_workbook(combined_path, data_only=False)
    for name in LAYOUT_SHEET_NAMES:
        if name in wb_formulas.sheetnames:
            del wb_formulas[name]
    wb_formulas.save(data_out_path)
    print(f"데이터 파일(데이터+매핑정보, 수식 유지) 저장 완료 -> {data_out_path}")
