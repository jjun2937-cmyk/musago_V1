"""
무사고전환 데이터/매핑정보 -> 결합 워크북(데이터+매핑정보+레이아웃4종) 생성.

이 스크립트가 만든 파일은 recalc.py로 재계산한 뒤, split_report_v4.py로
"데이터 파일"과 "결과 파일"(값만, 레이아웃 4종)로 분리해야 한다.
(결과 파일만 따로 저장하려면 레이아웃의 SUMPRODUCT 수식이 참조하는 데이터 시트가
그 파일 안에 없어야 하는데, 그러면 수식이 다 #REF! 로 깨지기 때문에 "재계산 후
값으로 굳혀서 분리"하는 2단계 방식을 쓴다.)

기존내용은 다 유지하고, 아래 4개 레이아웃을 만든다:
  레이아웃1_당월_부문별: 구분=수금부문명(개인/전략/신사업/법인, 포함매칭), 당월 지표만
  레이아웃2_전체_부문별: 위 + 기존전환완료/미완료(001~007 세분화 포함)
  레이아웃3_당월_상품별: 구분=상품명(데이터에 실제 존재하는 값, 정확매칭), 당월 지표만
  레이아웃4_전체_상품별: 위 + 기존전환완료/미완료(001~007 세분화 포함)

"기존전환완료/미완료" 정의:
  기존전환완료 = 사고有(당사, 사고구분코드='01') 중 현재판매플랜코드 != 최초판매플랜코드
                (사고전환판매플랜코드는 무사고년수=0이면 매핑정보에 값이 없어 항상 공란이
                 되므로, "최초판매플랜코드와 달라졌는지"로 전환 여부를 판단한다)
  기존전환미완료 = 사고有(당사) 중 현재판매플랜코드 == 최초판매플랜코드
                 (전환처리구분코드로 001~007 세분화)

"당월전환완료" 정의 (사고구분코드가 01/02든 공백이든 무관):
  현재판매플랜코드 == 경과전환판매플랜코드 인 계약 전체
"사고有(당사)" (당월 대상에서 빠지는 인원):
  사고구분코드='01' 이면서 현재판매플랜코드 != 경과전환판매플랜코드
  (01이어도 이미 경과전환판매플랜코드로 올라와 있으면 당월전환완료로 편입)

사용법: python build_report_v4.py <입력.xlsx> <결합출력.xlsx>
"""
import sys
import openpyxl
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# 헤더 탐색 유틸
# ---------------------------------------------------------------------------

def header_map(ws):
    m = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v is not None and v not in m:
            m[v] = c
    return m


def find_column(ws, hmap, name, fallback_anchor=None, fallback_expected=None, rename_to=None):
    if name in hmap:
        col = hmap[name]
    else:
        if fallback_anchor not in hmap:
            raise ValueError(f"'{name}' 헤더가 없고, 대체 기준 헤더 '{fallback_anchor}'도 찾을 수 없습니다.")
        anchor_col = hmap[fallback_anchor]
        col = anchor_col + 1
        actual = ws.cell(row=1, column=col).value
        if actual != fallback_expected:
            raise ValueError(
                f"'{fallback_anchor}' 바로 오른쪽 열의 헤더가 '{fallback_expected}'일 것으로 "
                f"예상했으나 실제로는 '{actual}' 입니다. (열 {col})"
            )
    if rename_to and ws.cell(row=1, column=col).value != rename_to:
        ws.cell(row=1, column=col).value = rename_to
    return col


def identify_sheets(wb):
    data_ws = mapping_ws = None
    for ws in wb.worksheets:
        hmap = header_map(ws)
        if '수금부문명' in hmap:
            data_ws = ws
        else:
            mapping_ws = ws
    if data_ws is None or mapping_ws is None:
        raise ValueError("데이터 시트('수금부문명' 헤더 보유) 또는 매핑정보 시트를 찾지 못했습니다.")
    return data_ws, mapping_ws


def find_or_append_column(ws, hmap, name):
    if name in hmap:
        return hmap[name]
    col = ws.max_column + 1
    ws.cell(row=1, column=col).value = name
    hmap[name] = col
    return col


# ---------------------------------------------------------------------------
# 1) 데이터 시트: 필요 컬럼 확정 + 사고전환판매플랜코드 채우기
# ---------------------------------------------------------------------------

def process_data_sheet(data_ws):
    hmap = header_map(data_ws)

    col_current = find_column(
        data_ws, hmap, '현재판매플랜코드',
        fallback_anchor='상품명', fallback_expected='판매플랜코드',
        rename_to='현재판매플랜코드',
    )
    col_initial = find_column(
        data_ws, hmap, '최초판매플랜코드',
        fallback_anchor='적용시작일시', fallback_expected='판매플랜코드',
        rename_to='최초판매플랜코드',
    )
    col_elapsed_plan = find_column(
        data_ws, hmap, '경과전환판매플랜코드',
        fallback_anchor='세부플랜코드', fallback_expected='전환판매플랜코드',
        rename_to='경과전환판매플랜코드',
    )
    hmap = header_map(data_ws)

    col_accident_free_years = hmap['무사고년수']
    col_accident_code = hmap['사고구분코드']
    col_process_code = hmap['전환처리구분코드']
    col_dept = hmap['수금부문명']
    col_product = hmap['상품명']
    col_month = hmap['계약체결월']

    col_T = find_or_append_column(data_ws, hmap, '사고전환판매플랜코드')

    return {
        'current': col_current, 'initial': col_initial, 'elapsed_plan': col_elapsed_plan,
        'accident_free_years': col_accident_free_years, 'accident_code': col_accident_code,
        'process_code': col_process_code, 'dept': col_dept, 'product': col_product,
        'month': col_month, 'T': col_T,
    }


def fill_data_formulas(data_ws, cols, mapping_sheet_name, map_key_col, map_target_col, map_last_row):
    L = get_column_letter
    K = L(cols['initial'])
    S = L(cols['accident_free_years'])
    map_key = L(map_key_col)
    map_target = L(map_target_col)

    first_row, last_row = 2, data_ws.max_row
    for r in range(first_row, last_row + 1):
        data_ws.cell(row=r, column=cols['T']).value = (
            f'=IFERROR(INDEX({mapping_sheet_name}!${map_target}$2:${map_target}${map_last_row},'
            f'MATCH({K}{r}&{S}{r},{mapping_sheet_name}!${map_key}$2:${map_key}${map_last_row},0)),"")'
        )


# ---------------------------------------------------------------------------
# 2) 매핑정보 시트: 키값 컬럼 생성
# ---------------------------------------------------------------------------

def process_mapping_sheet(mapping_ws):
    hmap = header_map(mapping_ws)
    col_plan = hmap['판매플랜코드']
    col_years = hmap['경과년수']
    col_target = hmap['전환판매플랜코드']
    col_key = find_or_append_column(mapping_ws, hmap, '키값')

    L = get_column_letter
    plan_l, years_l = L(col_plan), L(col_years)
    last_row = mapping_ws.max_row
    for r in range(2, last_row + 1):
        mapping_ws.cell(row=r, column=col_key).value = f'={plan_l}{r}&{years_l}{r}'

    return {'key': col_key, 'target': col_target, 'last_row': last_row}


def get_distinct_values(data_ws, col):
    vals = set()
    for r in range(2, data_ws.max_row + 1):
        v = data_ws.cell(row=r, column=col).value
        if v not in (None, ''):
            vals.add(str(v))
    return sorted(vals)


# ---------------------------------------------------------------------------
# 3) 결과레이아웃 시트 빌더 (구분 기준/기존 포함 여부에 따라 4종 생성)
# ---------------------------------------------------------------------------

DEPT_GROUPS = [('개인', '개인'), ('전략', '전략'), ('신사업', '신사업'), ('법인', '법인')]
CODES = [('001', '연락두절 등'), ('002', '사고有(타사)'), ('003', '전환의사無'),
         ('004', '압류계약'), ('005', '전환예정'), ('006', '기타'), ('007', 'ARS거부')]


def metric_headers(include_existing):
    headers = [
        '구분', '유지계약(A)', '사고有(당사)(B)', '당월전환대상(C=A-B)',
        '당월전환완료', '완료율(%)', '당월전환미완료', '미완료율(%)',
        '현장활동확인', '전체比(%)',
    ]
    headers += [f'{code} {label}' for code, label in CODES]
    headers += ['현장활동미확인', '전체比(%)']
    if include_existing:
        headers += ['기존전환완료', '유지계약比(%)', '기존전환미완료', '유지계약比(%)']
        headers += [f'{code} {label}(기존)' for code, label in CODES]
        headers += ['기존현장활동미확인', '유지계약比(%)']
    return headers


def build_one_layout(wb, data_ws, cols, months, group_col, group_defs, match_mode,
                      include_existing, sheet_name):
    """
    group_col: 'dept' 또는 'product' (cols 딕셔너리 키)
    group_defs: match_mode='contains' 이면 [(라벨,키워드),...], 'exact' 이면 [값,값,...]
    """
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    L = get_column_letter
    data_name = f"'{data_ws.title}'" if ' ' in data_ws.title else data_ws.title
    last_row = data_ws.max_row
    group_l = L(cols[group_col])
    month_l = L(cols['month'])
    v_l = L(cols['accident_code'])
    e_l = L(cols['current'])
    k_l = L(cols['initial'])
    m_l = L(cols['elapsed_plan'])
    u_l = L(cols['process_code'])

    headers = metric_headers(include_existing)
    # 중복 헤더('전체比(%)','유지계약比(%)')가 등장 순서대로 몇 번째 위치인지 기록
    positions = {}
    for i, h in enumerate(headers):
        positions.setdefault(h, []).append(i + 2)  # 2 = B열부터 시작

    for i, h in enumerate(headers):
        cell = ws.cell(row=1, column=i + 2)
        cell.value = h
        cell.font = openpyxl.styles.Font(bold=True)

    group_rng = f"{data_name}!${group_l}$2:${group_l}${last_row}"
    month_rng = f"{data_name}!${month_l}$2:${month_l}${last_row}"
    v_rng = f"{data_name}!${v_l}$2:${v_l}${last_row}"
    e_rng = f"{data_name}!${e_l}$2:${e_l}${last_row}"
    k_rng = f"{data_name}!${k_l}$2:${k_l}${last_row}"
    m_rng = f"{data_name}!${m_l}$2:${m_l}${last_row}"
    u_rng = f"{data_name}!${u_l}$2:${u_l}${last_row}"

    def C(name, r, occurrence=0):
        return f'{L(positions[name][occurrence])}{r}'

    def rng_of(name):
        col = L(positions[name][0])
        return col

    def write_row(r, month, extra_filter=None):
        base = f'({month_rng}="{month}")'
        if extra_filter:
            base += f'*{extra_filter}'

        ws[C('유지계약(A)', r)] = f'=SUMPRODUCT({base})'
        ws[C('사고有(당사)(B)', r)] = f'=SUMPRODUCT({base}*({v_rng}="01")*({e_rng}<>{m_rng}))'
        ws[C('당월전환대상(C=A-B)', r)] = f"={C('유지계약(A)', r)}-{C('사고有(당사)(B)', r)}"
        ws[C('당월전환완료', r)] = f'=SUMPRODUCT({base}*({e_rng}={m_rng}))'
        ws[C('완료율(%)', r)] = f"=IFERROR({C('당월전환완료', r)}/{C('당월전환대상(C=A-B)', r)}*100,\"\")"
        ws[C('당월전환미완료', r)] = f"={C('당월전환대상(C=A-B)', r)}-{C('당월전환완료', r)}"
        ws[C('미완료율(%)', r)] = f"=IFERROR({C('당월전환미완료', r)}/{C('당월전환대상(C=A-B)', r)}*100,\"\")"

        for code, label in CODES:
            name = f'{code} {label}'
            ws[C(name, r)] = f'=SUMPRODUCT({base}*({v_rng}<>"01")*({e_rng}<>{m_rng})*({u_rng}="{code}"))'
        first_code_col = L(positions[f'{CODES[0][0]} {CODES[0][1]}'][0])
        last_code_col = L(positions[f'{CODES[-1][0]} {CODES[-1][1]}'][0])
        ws[C('현장활동확인', r)] = f'=SUM({first_code_col}{r}:{last_code_col}{r})'
        ws[C('전체比(%)', r, 0)] = f"=IFERROR({C('현장활동확인', r)}/{C('당월전환대상(C=A-B)', r)}*100,\"\")"
        ws[C('현장활동미확인', r)] = f"={C('당월전환미완료', r)}-{C('현장활동확인', r)}"
        ws[C('전체比(%)', r, 1)] = f"=IFERROR({C('현장활동미확인', r)}/{C('당월전환대상(C=A-B)', r)}*100,\"\")"

        if include_existing:
            ws[C('기존전환완료', r)] = (
                f'=SUMPRODUCT({base}*({v_rng}="01")*({e_rng}<>{k_rng})*({e_rng}<>{m_rng}))'
            )
            ws[C('유지계약比(%)', r, 0)] = f"=IFERROR({C('기존전환완료', r)}/{C('유지계약(A)', r)}*100,\"\")"
            ws[C('기존전환미완료', r)] = f"={C('사고有(당사)(B)', r)}-{C('기존전환완료', r)}"
            ws[C('유지계약比(%)', r, 1)] = f"=IFERROR({C('기존전환미완료', r)}/{C('유지계약(A)', r)}*100,\"\")"
            for code, label in CODES:
                name = f'{code} {label}(기존)'
                ws[C(name, r)] = f'=SUMPRODUCT({base}*({v_rng}="01")*({e_rng}={k_rng})*({u_rng}="{code}"))'
            first_code2_col = L(positions[f'{CODES[0][0]} {CODES[0][1]}(기존)'][0])
            last_code2_col = L(positions[f'{CODES[-1][0]} {CODES[-1][1]}(기존)'][0])
            ws[C('기존현장활동미확인', r)] = (
                f"={C('기존전환미완료', r)}-SUM({first_code2_col}{r}:{last_code2_col}{r})"
            )
            ws[C('유지계약比(%)', r, 2)] = f"=IFERROR({C('기존현장활동미확인', r)}/{C('유지계약(A)', r)}*100,\"\")"

    blocks = []
    row = 2
    for month in months:
        total_row = row
        n_groups = len(group_defs)
        group_first = row + 1
        group_last = row + n_groups
        blocks.append((month, total_row, list(range(group_first, group_last + 1))))

        ws.cell(row=total_row, column=2).value = f"{month}월"
        ws.cell(row=total_row, column=2).font = openpyxl.styles.Font(bold=True)
        write_row(total_row, month)

        for gi, gdef in enumerate(group_defs):
            r = group_first + gi
            if match_mode == 'contains':
                label, keyword = gdef
                ws.cell(row=r, column=2).value = label
                write_row(r, month, extra_filter=f'ISNUMBER(SEARCH("{keyword}",{group_rng}))')
            else:
                ws.cell(row=r, column=2).value = gdef
                write_row(r, month, extra_filter=f'({group_rng}="{gdef}")')

        if match_mode == 'contains':
            unclassified_row = group_last + 1
            ws.cell(row=unclassified_row, column=2).value = '미분류(그룹매칭 안됨)'
            not_any_kw = ''.join(
                f'*(ISNUMBER(SEARCH("{kw}",{group_rng}))=FALSE)' for _, kw in group_defs
            )
            ws.cell(row=unclassified_row, column=3).value = f'=SUMPRODUCT(({month_rng}="{month}"){not_any_kw})'
            row = unclassified_row + 2
        else:
            row = group_last + 2

    return ws


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("Usage: python build_report_v4.py <입력.xlsx> <결합출력.xlsx>")
        sys.exit(1)

    in_path, out_path = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(in_path, data_only=False)
    data_ws, mapping_ws = identify_sheets(wb)
    print(f"데이터 시트 식별: '{data_ws.title}' ({data_ws.max_row}행 x {data_ws.max_column}열)")
    print(f"매핑정보 시트 식별: '{mapping_ws.title}' ({mapping_ws.max_row}행 x {mapping_ws.max_column}열)")

    map_info = process_mapping_sheet(mapping_ws)
    cols = process_data_sheet(data_ws)
    fill_data_formulas(
        data_ws, cols, mapping_ws.title,
        map_info['key'], map_info['target'], map_info['last_row'],
    )

    months = get_distinct_values(data_ws, cols['month'])
    products = get_distinct_values(data_ws, cols['product'])
    print("발견된 계약체결월:", months)
    print(f"발견된 상품명 ({len(products)}종):", products)

    build_one_layout(wb, data_ws, cols, months, 'dept', DEPT_GROUPS, 'contains', False, '레이아웃1_당월_부문별')
    build_one_layout(wb, data_ws, cols, months, 'dept', DEPT_GROUPS, 'contains', True, '레이아웃2_전체_부문별')
    build_one_layout(wb, data_ws, cols, months, 'product', products, 'exact', False, '레이아웃3_당월_상품별')
    build_one_layout(wb, data_ws, cols, months, 'product', products, 'exact', True, '레이아웃4_전체_상품별')

    wb.save(out_path)
    print(f"결합 파일(데이터+매핑정보+레이아웃4종) 저장 완료 -> {out_path}")
    print("다음 단계: recalc.py로 재계산 후 split_report_v4.py로 데이터/결과 파일 분리")
