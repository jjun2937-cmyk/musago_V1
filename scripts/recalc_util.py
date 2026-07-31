"""
LibreOffice(soffice)를 headless로 띄워 엑셀 수식을 전부 재계산하는 유틸리티.

openpyxl은 수식 문자열만 쓸 뿐 계산된 값(캐시)을 남기지 않으므로, 파일을 저장한
뒤 이 모듈로 한 번 재계산해야 값이 채워진다. LibreOffice(soffice)가 설치되어
있어야 한다 (Ubuntu/Debian: `apt-get install libreoffice-calc`, macOS:
`brew install --cask libreoffice`).
"""
import json
import subprocess
import sys
import tempfile
import time
from pathlib import Path

RECALCULATE_MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def _stamp(path):
    st = Path(path).stat()
    return st.st_mtime_ns, st.st_size


def _setup_macro(profile_dir: Path, timeout=30):
    url = profile_dir.as_uri()
    subprocess.run(
        ["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
        capture_output=True, timeout=timeout,
    )
    macro_dir = profile_dir / "user" / "basic" / "Standard"
    if not macro_dir.exists():
        raise RuntimeError("LibreOffice가 프로필을 생성하지 못했습니다.")
    (macro_dir / "Module1.xba").write_text(RECALCULATE_MACRO)
    return url


def recalc(filename, timeout=120):
    """엑셀 파일의 모든 수식을 재계산하고 파일에 값을 캐시로 저장한다.

    반환값: {"status": "success"|"errors_found", "total_errors": int,
             "error_cells": [...]} 또는 {"error": "..."}
    """
    abs_path = str(Path(filename).absolute())
    if not Path(abs_path).exists():
        return {"error": f"파일이 없습니다: {filename}"}

    with tempfile.TemporaryDirectory(prefix="recalc-lo-profile-") as profile_dir:
        try:
            profile_url = _setup_macro(Path(profile_dir), timeout=timeout)
        except FileNotFoundError:
            return {"error": "soffice(LibreOffice)를 찾을 수 없습니다. 설치가 필요합니다."}
        except subprocess.TimeoutExpired:
            return {"error": "LibreOffice 프로필 생성이 시간 초과되었습니다."}

        before = _stamp(abs_path)
        cmd = [
            "soffice", "--headless", "--norestore",
            f"-env:UserInstallation={profile_url}",
            "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
            abs_path,
        ]
        try:
            result = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            return {"error": f"재계산이 {timeout}초 내에 끝나지 않았습니다. timeout을 늘려서 재시도하세요."}

        if result.returncode != 0:
            detail = (result.stderr or "").strip() or f"soffice 종료코드 {result.returncode}"
            return {"error": f"LibreOffice 재계산 실패: {detail}"}
        if _stamp(abs_path) == before:
            return {"error": "LibreOffice가 종료됐지만 파일을 다시 쓰지 않았습니다. 다른 soffice 프로세스가 실행 중인지 확인하세요."}

    from openpyxl import load_workbook
    excel_errors = ["#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"]
    error_cells = []
    wb = load_workbook(filename, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if not hasattr(ws, "iter_rows"):
            continue
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(e in cell.value for e in excel_errors):
                    error_cells.append(f"{sheet_name}!{cell.coordinate}: {cell.value}")
    wb.close()

    return {
        "status": "success" if not error_cells else "errors_found",
        "total_errors": len(error_cells),
        "error_cells": error_cells[:50],
    }


def main():
    if len(sys.argv) < 2:
        print("Usage: python recalc_util.py <파일.xlsx> [timeout_seconds]")
        sys.exit(1)
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 120
    result = recalc(sys.argv[1], timeout=timeout)
    print(json.dumps(result, indent=2, ensure_ascii=False))
    sys.exit(1 if "error" in result else 0)


if __name__ == '__main__':
    main()
