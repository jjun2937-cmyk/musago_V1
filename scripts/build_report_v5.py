"""
무사고전환 데이터/매핑정보 -> 결합 워크북(데이터+매핑정보+레이아웃4종) 생성 (v5).

v4 대비 바뀐 점:
  1) "당월"/"기존" 모집단을 가르는 기준을 사고구분코드('01') 대신
     `경과년수 = 무사고년수` 여부로 바꿈.
     - 사고구분코드는 "이번 사이클에만" 찍히는 값이라(87.5%가 과거 사고 이력이
       있어도 사고구분코드는 공백), 과거 사고 이력 전체를 대변하지 못하는
       문제가 실사용 데이터에서 확인됨 (16,239건이 이미 전환 완료됐는데도
       "당월전환미완료"로 잘못 잡히던 문제).
  2) 기존전환완료 판정을 "사고전환판매플랜코드와 일치"가 아니라
     "최초판매플랜코드와 다른지"로 함 (사고전환판매플랜코드는 무사고년수=0이면
     매핑정보에 값이 없어 항상 공란이 되는 케이스가 많아서, 더 안정적인
     "최초판매플랜코드 대비 변경 여부"를 기준으로 삼음).
  3) 기존전환미완료도 전환처리구분코드(001~007)로 세분화.
  4) 결과 레이아웃에 사용자가 제공한 원본 템플릿의 서식(맑은 고딕 폰트, 헤더
     3행 병합, 회색/연두색 구분, 숫자서식, 열너비 등)을 그대로 재현.

지표 정의:
  유지계약(A)      = 그룹별 전체 계약 건수
  사고有(당사)(B)   = 경과년수 != 무사고년수 인 계약 전체 (= "기존" 모집단)
  당월전환대상       = 유지계약 - 사고有(당사) (= 경과년수 = 무사고년수 인 계약)
  당월전환완료       = 당월전환대상 중 현재판매플랜코드 = 경과전환판매플랜코드
  당월전환미완료     = 당월전환대상 - 당월전환완료 (전환처리구분코드 001~007 세분화)
  기존전환완료       = 사고有(당사) 중 현재판매플랜코드 != 최초판매플랜코드
  기존전환미완료     = 사고有(당사) - 기존전환완료 (전환처리구분코드 001~007 세분화)

사용법: python build_report_v5.py <입력.xlsx> <결합출력.xlsx>
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# 헤더 탐색 유틸
# ---------------------------------------------------------------------------

def header_map(ws):
    m = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v is not None and v not in m:
            m[v] = c
    return m


def find_column(ws, hmap, name, fallback_anchor=None, fallback_expected=None, rename_to=None):
    if name in hmap:
        col = hmap[name]
    else:
        if fallback_anchor not in hmap:
            raise ValueError(f"'{name}' 헤더가 없고, 대체 기준 헤더 '{fallback_anchor}'도 찾을 수 없습니다.")
        anchor_col = hmap[fallback_anchor]
        col = anchor_col + 1
        actual = ws.cell(row=1, column=col).value
        if actual != fallback_expected:
            raise ValueError(
                f"'{fallback_anchor}' 바로 오른쪽 열의 헤더가 '{fallback_expected}'일 것으로 "
                f"예상했으나 실제로는 '{actual}' 입니다. (열 {col})"
            )
    if rename_to and ws.cell(row=1, column=col).value != rename_to:
        ws.cell(row=1, column=col).value = rename_to
    return col


def identify_sheets(wb):
    data_ws = mapping_ws = None
    for ws in wb.worksheets:
        hmap = header_map(ws)
        if '수금부문명' in hmap:
            data_ws = ws
        else:
            mapping_ws = ws
    if data_ws is None or mapping_ws is None:
        raise ValueError("데이터 시트('수금부문명' 헤더 보유) 또는 매핑정보 시트를 찾지 못했습니다.")
    return data_ws, mapping_ws


def find_or_append_column(ws, hmap, name):
    if name in hmap:
        return hmap[name]
    col = ws.max_column + 1
    ws.cell(row=1, column=col).value = name
    hmap[name] = col
    return col


# ---------------------------------------------------------------------------
# 1) 데이터 시트: 필요 컬럼 확정 + 사고전환판매플랜코드 채우기 (참고용으로 계속 계산)
# ---------------------------------------------------------------------------

def process_data_sheet(data_ws):
    hmap = header_map(data_ws)

    col_current = find_column(
        data_ws, hmap, '현재판매플랜코드',
        fallback_anchor='상품명', fallback_expected='판매플랜코드',
        rename_to='현재판매플랜코드',
    )
    col_initial = find_column(
        data_ws, hmap, '최초판매플랜코드',
        fallback_anchor='적용시작일시', fallback_expected='판매플랜코드',
        rename_to='최초판매플랜코드',
    )
    col_elapsed_plan = find_column(
        data_ws, hmap, '경과전환판매플랜코드',
        fallback_anchor='세부플랜코드', fallback_expected='전환판매플랜코드',
        rename_to='경과전환판매플랜코드',
    )
    hmap = header_map(data_ws)

    col_elapsed_years = hmap['경과년수']
    col_accident_free_years = hmap['무사고년수']
    col_process_code = hmap['전환처리구분코드']
    col_dept = hmap['수금부문명']
    col_product = hmap['상품명']
    col_month = hmap['계약체결월']

    col_T = find_or_append_column(data_ws, hmap, '사고전환판매플랜코드')

    return {
        'current': col_current, 'initial': col_initial, 'elapsed_plan': col_elapsed_plan,
        'elapsed_years': col_elapsed_years, 'accident_free_years': col_accident_free_years,
        'process_code': col_process_code, 'dept': col_dept, 'product': col_product,
        'month': col_month, 'T': col_T,
    }


def fill_data_formulas(data_ws, cols, mapping_sheet_name, map_key_col, map_target_col, map_last_row):
    L = get_column_letter
    K = L(cols['initial'])
    S = L(cols['accident_free_years'])
    map_key = L(map_key_col)
    map_target = L(map_target_col)

    first_row, last_row = 2, data_ws.max_row
    for r in range(first_row, last_row + 1):
        data_ws.cell(row=r, column=cols['T']).value = (
            f'=IFERROR(INDEX({mapping_sheet_name}!${map_target}$2:${map_target}${map_last_row},'
            f'MATCH({K}{r}&{S}{r},{mapping_sheet_name}!${map_key}$2:${map_key}${map_last_row},0)),"")'
        )


# ---------------------------------------------------------------------------
# 2) 매핑정보 시트: 키값 컬럼 생성
# ---------------------------------------------------------------------------

def process_mapping_sheet(mapping_ws):
    hmap = header_map(mapping_ws)
    col_plan = hmap['판매플랜코드']
    col_years = hmap['경과년수']
    col_target = hmap['전환판매플랜코드']
    col_key = find_or_append_column(mapping_ws, hmap, '키값')

    L = get_column_letter
    plan_l, years_l = L(col_plan), L(col_years)
    last_row = mapping_ws.max_row
    for r in range(2, last_row + 1):
        mapping_ws.cell(row=r, column=col_key).value = f'={plan_l}{r}&{years_l}{r}'

    return {'key': col_key, 'target': col_target, 'last_row': last_row}


def get_distinct_values(data_ws, col):
    vals = set()
    for r in range(2, data_ws.max_row + 1):
        v = data_ws.cell(row=r, column=col).value
        if v not in (None, ''):
            vals.add(str(v))
    return sorted(vals)


# ---------------------------------------------------------------------------
# 3) 서식 (원본 템플릿 재현)
# ---------------------------------------------------------------------------

FONT_NAME = '맑은 고딕'
HEADER_FILL = PatternFill('solid', fgColor='F2F2F2')   # 원본 theme0/-5% 정확 재현(연회색)
SUB_FILL = PatternFill('solid', fgColor='E2EFDA')       # 원본 theme9/+80% 정확 재현(연두)
THIN = Side(style='thin')
MEDIUM = Side(style='medium')
COUNT_FMT = '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'
RATIO_FMT = '_-* #,##0.0_-;-* #,##0.0_-;_-* "-"_-;_-@_-'
L = get_column_letter

DEPT_GROUPS = [('개인', '개인'), ('전략', '전략'), ('신사업', '신사업'), ('법인', '법인')]
CODES = [('001', '연락두절 등'), ('002', '사고有(타사)'), ('003', '전환의사無'),
         ('004', '압류계약'), ('005', '전환예정'), ('006', '기타'), ('007', 'ARS거부')]


def style_header(cell, bold=True, fill=HEADER_FILL, left=None):
    cell.font = Font(name=FONT_NAME, size=10, bold=bold)
    cell.fill = fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = Border(left=left or THIN, right=THIN, top=THIN, bottom=THIN)


def style_data(cell, bold=False, number_format=COUNT_FMT):
    cell.font = Font(name=FONT_NAME, size=10, bold=bold)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    cell.number_format = number_format


def style_gubun(cell, align='right'):
    cell.font = Font(name=FONT_NAME, size=10, bold=True)
    cell.alignment = Alignment(horizontal=align, vertical='center')
    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    cell.number_format = COUNT_FMT


# ---------------------------------------------------------------------------
# 4) 결과레이아웃 시트 빌더
# ---------------------------------------------------------------------------

def col_layout(include_existing):
    """key(내부용) -> 순서. 'gubun'이 B열부터 시작."""
    keys = []
    keys += ['gubun', '유지계약', '사고유당사', '당월대상', '당월완료']
    keys += ['완료율', '당월미완료', '미완료율']
    keys += ['현장확인', '전체비_확인']
    keys += [f'당월_{c}' for c, _ in CODES]
    keys += ['현장미확인', '전체비_미확인']
    if include_existing:
        keys += ['기존완료', '유지비_완료', '기존미완료', '유지비_미완료']
        keys += ['기존현장확인', '유지비_기존확인']
        keys += [f'기존_{c}' for c, _ in CODES]
        keys += ['기존현장미확인', '유지비_기존미확인']
    return {key: i + 2 for i, key in enumerate(keys)}  # B열=2부터


HEADER_LABELS = {
    'gubun': '구분', '유지계약': '유지계약\nA', '사고유당사': '사고有(당사)\nB',
    '당월대상': '당월\n전환대상\nC (A-B)', '당월완료': '당월\n전환완료\nD',
    '완료율': '%', '당월미완료': '당월\n전환미완료\nE (C-D)', '미완료율': '%',
    '현장확인': '현장활동\n확인', '전체비_확인': '전체比',
    '현장미확인': '현장활동\n미확인', '전체비_미확인': '전체比',
    '기존완료': '기존\n전환완료', '유지비_완료': '유지계약比',
    '기존미완료': '기존\n전환미완료', '유지비_미완료': '유지계약比',
    '기존현장확인': '기존현장활동\n확인', '유지비_기존확인': '유지계약比',
    '기존현장미확인': '기존현장활동\n미확인', '유지비_기존미확인': '유지계약比',
}
for _code, _label in CODES:
    HEADER_LABELS[f'당월_{_code}'] = f'{_code} {_label}'
    HEADER_LABELS[f'기존_{_code}'] = f'{_code} {_label}'

RATIO_KEYS = {'완료율', '미완료율', '전체비_확인', '전체비_미확인',
              '유지비_완료', '유지비_미완료', '유지비_기존확인', '유지비_기존미확인'}

# 2~4행 3줄 병합 그룹 (medium 좌측 테두리로 구획 표시)
MERGE_3ROW = [
    ('gubun', None), ('유지계약', None), ('사고유당사', None), ('당월대상', None),
    ('당월완료', MEDIUM), ('당월미완료', None), ('기존완료', MEDIUM), ('기존미완료', MEDIUM),
]
# 3~4행 2줄 병합 그룹
MERGE_2ROW = ['현장확인', '현장미확인', '기존현장확인', '기존현장미확인']

ORIGINAL_WIDTHS = {
    'gubun': 9.25, '당월완료': 8.25, '완료율': 6.0, '당월미완료': 10.75, '미완료율': 6.0,
    '현장확인': 10.5, '전체비_확인': 6.5, '당월_001': 10.125,
    '기존현장확인': 12.5, '기존현장미확인': 12.5,
}


def write_header(ws, col, include_existing):
    for key, left in MERGE_3ROW:
        if key not in col:
            continue
        ci = col[key]
        ws.merge_cells(start_row=2, start_column=ci, end_row=4, end_column=ci)
        for r in range(2, 5):
            style_header(ws.cell(row=r, column=ci), bold=True, fill=HEADER_FILL, left=left)
        ws.cell(row=2, column=ci, value=HEADER_LABELS[key])

    for key in MERGE_2ROW:
        if key not in col:
            continue
        ci = col[key]
        ws.merge_cells(start_row=3, start_column=ci, end_row=4, end_column=ci)
        for r in range(3, 5):
            style_header(ws.cell(row=r, column=ci), bold=True, fill=HEADER_FILL)
        ws.cell(row=3, column=ci, value=HEADER_LABELS[key])

    row4_only = ['완료율', '미완료율', '전체비_확인', '전체비_미확인',
                 '유지비_완료', '유지비_미완료', '유지비_기존확인', '유지비_기존미확인']
    for key in row4_only:
        if key not in col:
            continue
        cell = ws.cell(row=4, column=col[key], value=HEADER_LABELS[key])
        style_header(cell, bold=False, fill=SUB_FILL)

    for code, _ in CODES:
        for prefix in ('당월_', '기존_'):
            key = f'{prefix}{code}'
            if key not in col:
                continue
            cell = ws.cell(row=4, column=col[key], value=HEADER_LABELS[key])
            style_header(cell, bold=False, fill=HEADER_FILL)

    for key, width in ORIGINAL_WIDTHS.items():
        if key in col:
            ws.column_dimensions[L(col[key])].width = width
    ws.row_dimensions[1].height = 17.25
    ws.row_dimensions[2].height = 16.5
    ws.row_dimensions[3].height = 16.5


def build_one_layout(wb, data_ws, cols, months, group_col, group_defs, match_mode,
                      include_existing, sheet_name):
    """
    group_col: 'dept' 또는 'product' (cols 딕셔너리 키)
    group_defs: match_mode='contains' 이면 [(라벨,키워드),...], 'exact' 이면 [값,값,...]
    """
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    data_name = f"'{data_ws.title}'" if ' ' in data_ws.title else data_ws.title
    last_row = data_ws.max_row
    group_l = L(cols[group_col])
    month_l = L(cols['month'])
    e_l = L(cols['current'])
    k_l = L(cols['initial'])
    m_l = L(cols['elapsed_plan'])
    ey_l = L(cols['elapsed_years'])
    afy_l = L(cols['accident_free_years'])
    u_l = L(cols['process_code'])

    col = col_layout(include_existing)
    write_header(ws, col, include_existing)

    group_rng = f"{data_name}!${group_l}$2:${group_l}${last_row}"
    month_rng = f"{data_name}!${month_l}$2:${month_l}${last_row}"
    e_rng = f"{data_name}!${e_l}$2:${e_l}${last_row}"
    k_rng = f"{data_name}!${k_l}$2:${k_l}${last_row}"
    m_rng = f"{data_name}!${m_l}$2:${m_l}${last_row}"
    ey_rng = f"{data_name}!${ey_l}$2:${ey_l}${last_row}"
    afy_rng = f"{data_name}!${afy_l}$2:${afy_l}${last_row}"
    u_rng = f"{data_name}!${u_l}$2:${u_l}${last_row}"

    def C(key, r):
        return f'{L(col[key])}{r}'

    def write_row(r, month, extra_filter=None, bold=False):
        base = f'({month_rng}="{month}")'
        if extra_filter:
            base += f'*{extra_filter}'

        ws[C('유지계약', r)] = f'=SUMPRODUCT({base})'
        ws[C('사고유당사', r)] = f'=SUMPRODUCT({base}*({ey_rng}<>{afy_rng}))'
        ws[C('당월대상', r)] = f"={C('유지계약', r)}-{C('사고유당사', r)}"
        ws[C('당월완료', r)] = f'=SUMPRODUCT({base}*({ey_rng}={afy_rng})*({e_rng}={m_rng}))'
        ws[C('완료율', r)] = f"=IFERROR({C('당월완료', r)}/{C('당월대상', r)}*100,\"\")"
        ws[C('당월미완료', r)] = f"={C('당월대상', r)}-{C('당월완료', r)}"
        ws[C('미완료율', r)] = f"=IFERROR({C('당월미완료', r)}/{C('당월대상', r)}*100,\"\")"

        for code, _ in CODES:
            ws[C(f'당월_{code}', r)] = (
                f'=SUMPRODUCT({base}*({ey_rng}={afy_rng})*({e_rng}<>{m_rng})*({u_rng}="{code}"))'
            )
        first_c = C(f'당월_{CODES[0][0]}', r)
        last_c = C(f'당월_{CODES[-1][0]}', r)
        ws[C('현장확인', r)] = f'=SUM({first_c}:{last_c})'
        ws[C('전체비_확인', r)] = f"=IFERROR({C('현장확인', r)}/{C('당월대상', r)}*100,\"\")"
        ws[C('현장미확인', r)] = f"={C('당월미완료', r)}-{C('현장확인', r)}"
        ws[C('전체비_미확인', r)] = f"=IFERROR({C('현장미확인', r)}/{C('당월대상', r)}*100,\"\")"

        if include_existing:
            ws[C('기존완료', r)] = f'=SUMPRODUCT({base}*({ey_rng}<>{afy_rng})*({e_rng}<>{k_rng}))'
            ws[C('유지비_완료', r)] = f"=IFERROR({C('기존완료', r)}/{C('유지계약', r)}*100,\"\")"
            ws[C('기존미완료', r)] = f"={C('사고유당사', r)}-{C('기존완료', r)}"
            ws[C('유지비_미완료', r)] = f"=IFERROR({C('기존미완료', r)}/{C('유지계약', r)}*100,\"\")"
            for code, _ in CODES:
                ws[C(f'기존_{code}', r)] = (
                    f'=SUMPRODUCT({base}*({ey_rng}<>{afy_rng})*({e_rng}={k_rng})*({u_rng}="{code}"))'
                )
            first_c2 = C(f'기존_{CODES[0][0]}', r)
            last_c2 = C(f'기존_{CODES[-1][0]}', r)
            ws[C('기존현장확인', r)] = f'=SUM({first_c2}:{last_c2})'
            ws[C('유지비_기존확인', r)] = f"=IFERROR({C('기존현장확인', r)}/{C('유지계약', r)}*100,\"\")"
            ws[C('기존현장미확인', r)] = f"={C('기존미완료', r)}-{C('기존현장확인', r)}"
            ws[C('유지비_기존미확인', r)] = f"=IFERROR({C('기존현장미확인', r)}/{C('유지계약', r)}*100,\"\")"

        for key, ci in col.items():
            if key == 'gubun':
                continue
            style_data(ws.cell(row=r, column=ci), bold=bold,
                       number_format=RATIO_FMT if key in RATIO_KEYS else COUNT_FMT)

    blocks = []
    row = 5
    for month in months:
        total_row = row
        n_groups = len(group_defs)
        group_first = row + 1
        group_last = row + n_groups
        blocks.append((month, total_row, list(range(group_first, group_last + 1))))

        ws.cell(row=total_row, column=col['gubun'], value=f"{month}월")
        style_gubun(ws.cell(row=total_row, column=col['gubun']), align='left')
        write_row(total_row, month, bold=False)

        for gi, gdef in enumerate(group_defs):
            r = group_first + gi
            if match_mode == 'contains':
                label, keyword = gdef
                ws.cell(row=r, column=col['gubun'], value=label)
                style_gubun(ws.cell(row=r, column=col['gubun']), align='right')
                write_row(r, month, extra_filter=f'ISNUMBER(SEARCH("{keyword}",{group_rng}))')
            else:
                ws.cell(row=r, column=col['gubun'], value=gdef)
                style_gubun(ws.cell(row=r, column=col['gubun']), align='right')
                write_row(r, month, extra_filter=f'({group_rng}="{gdef}")')

        if match_mode == 'contains':
            unclassified_row = group_last + 1
            ws.cell(row=unclassified_row, column=col['gubun'], value='미분류(그룹매칭 안됨)')
            style_data(ws.cell(row=unclassified_row, column=col['gubun']), number_format='General')
            not_any_kw = ''.join(
                f'*(ISNUMBER(SEARCH("{kw}",{group_rng}))=FALSE)' for _, kw in group_defs
            )
            ws.cell(row=unclassified_row, column=col['유지계약'],
                    value=f'=SUMPRODUCT(({month_rng}="{month}"){not_any_kw})')
            style_data(ws.cell(row=unclassified_row, column=col['유지계약']))
            row = unclassified_row + 2
        else:
            row = group_last + 2

    return ws


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("Usage: python build_report_v5.py <입력.xlsx> <결합출력.xlsx>")
        sys.exit(1)

    in_path, out_path = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(in_path, data_only=False)
    data_ws, mapping_ws = identify_sheets(wb)
    print(f"데이터 시트 식별: '{data_ws.title}' ({data_ws.max_row}행 x {data_ws.max_column}열)")
    print(f"매핑정보 시트 식별: '{mapping_ws.title}' ({mapping_ws.max_row}행 x {mapping_ws.max_column}열)")

    map_info = process_mapping_sheet(mapping_ws)
    cols = process_data_sheet(data_ws)
    fill_data_formulas(
        data_ws, cols, mapping_ws.title,
        map_info['key'], map_info['target'], map_info['last_row'],
    )

    months = get_distinct_values(data_ws, cols['month'])
    products = get_distinct_values(data_ws, cols['product'])
    print("발견된 계약체결월:", months)
    print(f"발견된 상품명 ({len(products)}종):", products)

    build_one_layout(wb, data_ws, cols, months, 'dept', DEPT_GROUPS, 'contains', False, '레이아웃1_당월_부문별')
    build_one_layout(wb, data_ws, cols, months, 'dept', DEPT_GROUPS, 'contains', True, '레이아웃2_전체_부문별')
    build_one_layout(wb, data_ws, cols, months, 'product', products, 'exact', False, '레이아웃3_당월_상품별')
    build_one_layout(wb, data_ws, cols, months, 'product', products, 'exact', True, '레이아웃4_전체_상품별')

    wb.save(out_path)
    print(f"결합 파일(데이터+매핑정보+레이아웃4종) 저장 완료 -> {out_path}")
    print("다음 단계: recalc.py로 재계산 후 split_report_v5.py로 데이터/결과 파일 분리")
